from decimal import Decimal
import io
import uuid

from django.http import HttpResponse
from django.utils import timezone
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt
from khayyam import JalaliDatetime

from apps.order.enums import OrderTypeEnum
from apps.order.models import InvoiceModel, OrderModel
from apps.product.enums import CoinCategoryEnum


class InvoiceService:

    @staticmethod
    def generate_invoice_number():
        return f"INV-{timezone.now().strftime('%Y%m%d')}-{uuid.uuid4().hex[:6].upper()}"

    @staticmethod
    def calculate_weight(order):
        weight = Decimal("0")

        for item in order.items.all():
            if item.product:
                weight += item.product.weight or 0

            if item.coin:
                weight += item.coin.weight or 0

            if item.gold_amount > 0:
                weight += item.gold_amount

        return weight

    @staticmethod
    def calculate_tax(order):
        # فعلاً همه بدون مالیات
        return Decimal("0")

    @staticmethod
    def calculate_wage(order):
        # فعلاً همه بدون اجرت
        return Decimal("0")

    @staticmethod
    def generate_invoice(order: OrderModel):
        if hasattr(order, "invoice"):
            return order.invoice

        weight = InvoiceService.calculate_weight(order)
        tax_amount = InvoiceService.calculate_tax(order)
        wage_amount = InvoiceService.calculate_wage(order)

        # عنوان فاکتور بر اساس نوع سفارش
        if order.order_type == OrderTypeEnum.BUY_MELTED_GOLD:
            invoice_title = "فاکتور خرید طلای آب‌شده"
        elif order.order_type == OrderTypeEnum.BUY_PRODUCT:
            invoice_title = "فاکتور خرید محصول زینتی"
        elif order.order_type == OrderTypeEnum.BUY_COIN:
            invoice_title = "فاکتور خرید سکه"
        else:
            invoice_title = "فاکتور خرید"

        invoice = InvoiceModel.objects.create(
            order=order,
            invoice_number=InvoiceService.generate_invoice_number(),
            title=invoice_title,
            amount=order.total_amount,
            tax_amount=tax_amount,
            wage_amount=wage_amount,
            is_paid=True,
            paid_at=timezone.now(),
            weight=weight,
        )

        order.invoice_generated = True
        order.save()

        return invoice


class InvoiceDOCXService:

    @staticmethod
    def to_jalali(dt):
        if not dt:
            return "-"
        local_time = timezone.localtime(dt)
        return JalaliDatetime(local_time).strftime("%Y/%m/%d %H:%M")

    @staticmethod
    def generate_docx(invoice: InvoiceModel):
        doc = Document()

        # تنظیم فونت
        style = doc.styles["Normal"]
        style.font.name = "B Nazanin"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "B Nazanin")
        style.font.size = Pt(13)

        section = doc.sections[0]

        # هدر
        header = section.header
        header_p = header.paragraphs[0]
        header_p.text = "KIANMEHR"
        header_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        doc.add_paragraph("")

        # عنوان فاکتور
        title = doc.add_heading(level=1)
        run = title.add_run(f"{invoice.title} ({invoice.invoice_number})")
        run.font.name = "B Nazanin"

        # مشخصات مشتری
        # مشخصات مشتری
        doc.add_heading("مشخصات مشتری", level=2)

        table = doc.add_table(rows=2, cols=2)
        table.style = "Table Grid"

        table.rows[0].cells[0].text = "نام مشتری"
        table.rows[0].cells[1].text = invoice.order.user.full_name() or "-"

        table.rows[1].cells[0].text = "شماره تماس"
        table.rows[1].cells[1].text = invoice.order.user.phone_number

        # مشخصات فاکتور
        doc.add_heading("مشخصات فاکتور", level=2)

        final_amount = invoice.amount + invoice.tax_amount + invoice.wage_amount

        data = [
            ("مبلغ کل", f"{invoice.amount:,} تومان"),
            ("مبلغ نهایی فاکتور", f"{final_amount:,} تومان"),
            ("وضعیت پرداخت", "پرداخت شده" if invoice.is_paid else "پرداخت نشده"),
            ("تاریخ پرداخت", InvoiceDOCXService.to_jalali(invoice.paid_at)),
            (
                "زمان قفل قیمت",
                InvoiceDOCXService.to_jalali(invoice.order.locked_price_at),
            ),
        ]

        if invoice.order.order_type == OrderTypeEnum.BUY_MELTED_GOLD:
            data.append(("وزن کل", f"{invoice.weight} گرم"))

        # سکه
        if invoice.order.order_type == OrderTypeEnum.BUY_COIN:
            item = invoice.order.items.first()
            if item.coin.category == CoinCategoryEnum.PARSIAN:
                data.append(("نوع سکه", f"پارسیان {item.coin.weight} گرم"))
            else:
                data.append(("نوع سکه", item.coin.coin_type))
            data.append(("دسته‌بندی سکه", item.coin.category))

        table2 = doc.add_table(rows=len(data), cols=2)
        table2.style = "Table Grid"
        for i, (label, value) in enumerate(data):
            table2.rows[i].cells[0].text = label
            table2.rows[i].cells[1].text = value

        # آیتم‌های سفارش
        doc.add_heading("آیتم‌های سفارش", level=2)
        items = invoice.order.items.all()

        table3 = doc.add_table(rows=1, cols=5)
        table3.style = "Table Grid"
        hdr = table3.rows[0].cells
        hdr[0].text = "نوع"
        hdr[1].text = "عنوان"
        hdr[2].text = "وزن / تعداد"
        hdr[3].text = "قیمت واحد"
        hdr[4].text = "قیمت کل"

        for item in items:
            row = table3.add_row().cells

            if item.product:
                row[0].text = "محصول زینتی"
                row[1].text = f"{item.product.title} (ID: {item.product.id})"
                row[2].text = f"{item.product.weight} گرم"

            elif item.coin:
                row[0].text = "سکه"

                if item.coin.category == CoinCategoryEnum.PARSIAN:
                    row[1].text = f"سکه پارسیان {item.coin.weight} گرم"
                else:
                    row[1].text = f"سکه {item.coin.coin_type}"

                row[2].text = f"{item.quantity} عدد"

            elif item.gold_amount > 0:
                row[0].text = "طلای آب‌شده"
                row[1].text = "خرید طلای آب‌شده"
                row[2].text = f"{item.gold_amount} گرم"

            else:
                row[0].text = "-"
                row[1].text = "-"
                row[2].text = "-"

            row[3].text = f"{item.unit_price:,}"
            row[4].text = f"{item.total_price:,}"

        buffer = io.BytesIO()
        doc.save(buffer)
        buffer.seek(0)

        response = HttpResponse(
            buffer.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        response["Content-Disposition"] = (
            f"attachment; filename={invoice.invoice_number}.docx"
        )

        return response
